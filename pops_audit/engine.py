"""Compare confirmed business cells, never the entire visual worksheet grid."""
from collections import Counter, defaultdict
from datetime import datetime, timezone
from fnmatch import fnmatchcase
import hashlib
import json
import math
from pathlib import Path
import unicodedata

from openpyxl.utils.cell import get_column_letter, range_boundaries
from .formulas import normalize, same_formula, has_broken_reference
from .reader import read_workbook
from .locator import locate


CATEGORIES = {
    "FILE_MISSING": ("File not received", "medium"),
    "FILE_UNEXPECTED": ("Unmatched received file", "medium"),
    "FILE_AMBIGUOUS": ("Ambiguous file pairing", "medium"),
    "FILE_READ_ERROR": ("Workbook could not be read", "medium"),
    "SHEET_MISSING": ("Sheet removed", "medium"),
    "SHEET_ADDED": ("Sheet added", "medium"),
    "SHEET_RENAMED": ("Sheet name changed", "medium"),
    "KPI_INTEGRITY": ("KPI integrity", "medium"),
    "KPI_MISSING": ("Required KPI missing", "high"),
    "KPI_UNEXPECTED": ("Unexpected KPI occurrence", "medium"),
    "COLUMN_MISSING": ("Required column missing", "medium"),
    "COLUMN_ADDED": ("Table column added", "medium"),
    "LAYOUT_AMBIGUOUS": ("Comparison needs a layout check", "medium"),
    "LAYOUT_CHANGED": ("KPI or column position changed", "medium"),
    "FORMULA_REPLACED": ("Formula replaced with a value", "medium"),
    "FORMULA_MODIFIED": ("Formula modified", "medium"),
    "FORMULA_REMOVED": ("Formula removed", "medium"),
    "VALUE_TO_FORMULA": ("Prefilled value replaced with a formula", "medium"),
    "PREFILLED_VALUE_CHANGED": ("Prefilled value changed", "medium"),
    "BROKEN_REFERENCE": ("Broken reference", "medium"),
    "CELL_ERROR": ("Excel calculation error", "medium"),
    "INVALID_INPUT": ("Invalid numeric input", "low"),
    "TEMPLATE_ISSUE": ("Sent template or configuration issue", "medium"),
    "EXTERNAL_LINK_CHANGED": ("External workbook links changed", "medium"),
    "DEFINED_NAME_CHANGED": ("Referenced workbook name changed", "medium"),
}


def fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     default=str, allow_nan=False).encode("utf-8")).hexdigest()[:24]


def _json_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def snapshot(cell):
    if cell is None:
        return None
    return {"value": _json_value(cell.value), "formula": cell.formula,
            "kind": cell.kind, "cache_available": cell.cache_available}


def blank(cell):
    return cell is None or not cell.formula and (cell.value is None or isinstance(cell.value, str) and not cell.value.strip())


def numeric(value):
    return not isinstance(value, bool) and isinstance(value, (int, float)) and math.isfinite(value)


def state(cell):
    if cell is not None and cell.formula and not cell.cache_available:
        return "unknown"
    if cell is not None and cell.kind not in {"b", "e", "d"} and numeric(cell.value) and cell.value != 0:
        return "filled"
    return "invalid"


def filling(counts=None):
    counts = counts or {}
    result = {key: int(counts.get(key, 0)) for key in ("filled", "invalid", "unknown")}
    expected = sum(result.values())
    known = expected - result["unknown"]
    result.update(expected=expected, rate=result["filled"] / expected if expected and not result["unknown"] else None,
                  known_rate=result["filled"] / known if known else None,
                  coverage=known / expected if expected else None)
    return result


def sum_filling(items):
    return filling({key: sum(item.get(key, 0) for item in items) for key in ("filled", "invalid", "unknown")})


class Findings:
    def __init__(self, config, country):
        self.config, self.country, self.items = config, country, []
        self.seen = set()

    def add(self, category, message, *, sheet="", kpi="", column="", sent_cell="", received_cell="",
            before=None, after=None, baseline=False, title=None, severity=None):
        if self.config.get("checks", {}).get(category, True) is False:
            return
        label, default_severity = CATEGORIES[category]
        finding = {"country_id": self.country["id"], "country": self.country["name"],
                   "workbook": self.country.get("received_file") or self.country.get("sent_file") or self.country["file_key"],
                   "sheet": sheet, "category": category, "category_label": label,
                   "severity": severity or default_severity,
                   "title": title or label, "message": message, "kpi": kpi, "column": column,
                   "sent_cell": sent_cell, "received_cell": received_cell,
                   "before": before, "after": after, "baseline": baseline, "status": "unreviewed"}
        identifier = fingerprint(finding)
        if identifier not in self.seen:
            finding["id"] = identifier
            self.items.append(finding)
            self.seen.add(identifier)


def _same_value(before, after, config):
    if numeric(before) and numeric(after):
        return math.isclose(before, after, rel_tol=config.get("relative_tolerance", 1e-9),
                            abs_tol=config.get("absolute_tolerance", 1e-9))
    return type(before) is type(after) and before == after


def _cell_checks(before, after, context, findings, config, rule, references, sent_name, received_name):
    if after is None:
        return  # A structural finding already explains this missing location.
    kwargs = {**context, "before": snapshot(before), "after": snapshot(after)}
    broken_before = has_broken_reference(before.formula) or before.kind == "e" and before.value == "#REF!"
    broken_after = has_broken_reference(after.formula) or after.kind == "e" and after.value == "#REF!"
    if broken_after:
        unchanged_broken = broken_before and (
            bool(before.formula) == bool(after.formula)
            and (same_formula(before.formula, after.formula, sent_name, received_name, references,
                              ignore_case=config.get("ignore_formula_case", True))
                 if before.formula else _same_value(before.value, after.value, config)))
        findings.add("BROKEN_REFERENCE", "The received cell contains a broken #REF! reference."
                     if not unchanged_broken else "This broken #REF! reference was already present in the sent template.",
                     baseline=unchanged_broken, **kwargs)
        return
    if before.formula:
        if not after.formula:
            findings.add("FORMULA_REMOVED" if blank(after) else "FORMULA_REPLACED",
                         "A formula supplied in the sent template has been cleared." if blank(after) else
                         "A formula supplied in the sent template has been replaced by a constant; its automatic link is lost.", **kwargs)
            return
        if not same_formula(before.formula, after.formula, sent_name, received_name, references,
                            ignore_case=config.get("ignore_formula_case", True)):
            findings.add("FORMULA_MODIFIED", "The received formula differs from the sent formula after normalising harmless syntax and confirmed moved references.", **kwargs)
            return
    elif not blank(before) and after.formula:
        findings.add("VALUE_TO_FORMULA", "A prefilled constant has become a formula. Check that this change is allowed.", **kwargs)
        return
    elif not blank(before) and not after.formula:
        historical = {normalize(value) for value in rule.get("historical_columns", config.get("historical_columns", []))}
        if (config.get("prefilled_changes", "historical") == "all" or normalize(context["column"]) in historical):
            if not _same_value(before.value, after.value, config):
                findings.add("PREFILLED_VALUE_CHANGED", "A prefilled historical value has changed or been cleared."
                             if normalize(context["column"]) in historical else "A prefilled value has changed or been cleared.", **kwargs)
                return
    if after.kind == "e":
        existing = before.kind == "e" and before.value == after.value
        findings.add("CELL_ERROR", "Excel saved a calculation error for this cell. The error was already present in the sent template."
                     if existing else "Excel saved a calculation error for this received cell. Check the formula and its inputs.", baseline=existing, **kwargs)
    elif (after.cache_available and after.value is not None
          and not (isinstance(after.value, str) and not after.value.strip()) and state(after) == "invalid"):
        # Unchanged template placeholders/zeros are reflected in filling, without
        # creating thousands of fictitious country changes.
        if (not _same_value(before.value, after.value, config) or before.kind != after.kind
                or state(before) == "filled"):
            if isinstance(after.value, str):
                message = ("The saved formula result is text, not a number. Check the formula's inputs; keep its automatic link intact."
                           if after.formula else
                           "Excel stores this entry as text, not as a number. A number displayed as 0,25 or 25% is valid; "
                           "text such as '0,25' is not usable by the intermediary-table extractor. Convert this cell to a number in Excel.")
            elif isinstance(after.value, bool):
                message = "This entry is a Boolean (TRUE/FALSE). A nonzero number or numeric percentage is required."
            elif numeric(after.value) and after.value == 0:
                message = "This entry is zero. A nonzero number or numeric percentage is required to count as filled."
            else:
                message = "This entry is not a finite number. A nonzero number or numeric percentage is required."
            findings.add("INVALID_INPUT", message, **kwargs)


def _ignored(rule, row, column):
    for reference in rule.get("ignore_ranges", []):
        left, top, right, bottom = range_boundaries(reference)
        if left <= column <= right and top <= row <= bottom:
            return True
    return False


def _name_group(row, mode):
    if mode == "layout" and isinstance(row.key, (list, tuple)):
        return str(row.key[0])
    if mode == "identifiers":
        # Locator keys already canonicalise numeric representations. Display
        # strings such as '1' and '1.0' must not create two identifier groups.
        return tuple(row.key) if isinstance(row.key, list) else row.key
    return normalize(row.identifier)


def _groups(layout, mode):
    groups = defaultdict(list)
    if layout:
        for row in layout.rows:
            groups[_name_group(row, mode)].append(row)
    return groups


def _row_key(row):
    return tuple(row.key) if isinstance(row.key, list) else row.key


def _columns(layout):
    return sorted(set(getattr(layout, "columns", []) or []).union(
        *(set(row.columns) for row in layout.rows))) if layout else []


def _resolve_sheet(book, rule):
    allowed = {normalize(value) for value in [rule["name"], *rule.get("aliases", [])]}
    matches = [sheet for name, sheet in book.sheets.items() if normalize(name) in allowed]
    return matches[0] if len(matches) == 1 else None, matches


def _problems(layout, findings, name, *, baseline=False, sent_layout=None, mode="layout"):
    if not layout:
        return
    sent_problems = {(p["code"], p.get("identifier"), p.get("column")) for p in sent_layout.problems} if sent_layout else set()
    pending, column_groups = [], {}
    for problem in layout.problems:
        code = problem["code"]
        if mode == "identifiers" and code == "DUPLICATE_IDENTIFIER":
            continue  # One grouped integrity finding below, with identifiers.
        if code in {"MISSING_KPI", "EXTRA_KPI"} and not baseline:
            continue  # Sent/received counts produce a single structural finding.
        already_present = (code, problem.get("identifier"), problem.get("column")) in sent_problems
        if not baseline and already_present:
            continue
        if code in {"MISSING_COLUMN", "AMBIGUOUS_COLUMN"}:
            # One absent period header can affect hundreds of KPI rows. Keep
            # the evidence together instead of repeating the same action.
            key = code, problem.get("column")
            if key not in column_groups:
                column_groups[key] = {**problem, "coordinates": [], "identifiers": []}
                pending.append(column_groups[key])
            grouped = column_groups[key]
            grouped["coordinates"].extend(problem.get("coordinates", []))
            if problem.get("identifier"):
                grouped["identifiers"].append(str(problem["identifier"]))
                grouped["coordinates"].extend(
                    f"{get_column_letter(row.label_column)}{row.row}" for row in layout.rows
                    if normalize(row.identifier) == normalize(problem["identifier"]))
        else:
            pending.append(problem)
    for problem in pending:
        code = problem["code"]
        category = "TEMPLATE_ISSUE" if baseline else (
            "KPI_INTEGRITY" if code in {"MISSING_IDENTIFIER", "DUPLICATE_IDENTIFIER"} else
            "COLUMN_MISSING" if code == "MISSING_COLUMN" else "LAYOUT_AMBIGUOUS")
        coordinates = ", ".join(dict.fromkeys(problem.get("coordinates", [])))
        identifiers = list(dict.fromkeys(problem.get("identifiers", [])))
        message = problem["message"]
        if len(identifiers) > 1:
            message += " Affected KPI labels: " + ", ".join(identifiers) + "."
        findings.add(category, message, sheet=name,
                     kpi=identifiers[0] if len(identifiers) == 1 else "" if identifiers else problem.get("identifier", ""),
                     column=problem.get("column", ""), baseline=baseline,
                     sent_cell=coordinates if baseline else "", received_cell="" if baseline else coordinates)


def _integrity(before, after, rule, findings, name):
    mode = rule["mode"]
    first, second = _groups(before, mode), _groups(after, mode)
    ambiguous = set()
    for key in first.keys() | second.keys():
        old, new = first.get(key, []), second.get(key, [])
        label = (old or new)[0].identifier
        if mode == "identifiers" and (len(old) > 1 or len(new) > 1):
            ambiguous.add(key)
            findings.add("KPI_INTEGRITY", f"Identifier {label!r} occurs {len(old)} time(s) in the sent table and {len(new)} time(s) in the received table. An identifier must be unique; affected values were not paired.",
                         sheet=name, kpi=label, baseline=len(old) > 1 and len(new) == len(old),
                         sent_cell=", ".join(f"{get_column_letter(row.label_column)}{row.row}" for row in old),
                         received_cell=", ".join(f"{get_column_letter(row.label_column)}{row.row}" for row in new))
        elif len(old) != len(new):
            if mode == "identifiers" and ((not new and _unresolved_layout(after)) or (not old and _unresolved_layout(before))):
                # An unreadable identifier may be the supposedly missing row.
                # The locator's ambiguity finding already asks for review.
                ambiguous.add(key)
                continue
            if mode == "layout" and old and new:
                ambiguous.add(key)
            category = "KPI_INTEGRITY" if mode == "identifiers" else "KPI_MISSING" if len(new) < len(old) else "KPI_UNEXPECTED"
            findings.add(category, f"{label!r}: {len(old)} occurrence(s) in the sent file, {len(new)} in the received file. " +
                         ("Repeated names with unequal counts are not paired by order." if key in ambiguous else "Review the missing or additional identifier."),
                         sheet=name, kpi=label,
                         sent_cell=", ".join(f"{get_column_letter(row.label_column)}{row.row}" for row in old),
                         received_cell=", ".join(f"{get_column_letter(row.label_column)}{row.row}" for row in new),
                         before={"value": f"{len(old)} occurrence(s)", "formula": None},
                         after={"value": f"{len(new)} occurrence(s)", "formula": None})
        if any(row.ambiguous for row in old + new):
            ambiguous.add(key)
    return first, second, ambiguous


def _prepare_sheets(before_book, after_book, config, findings):
    prepared = []
    claimed_before, claimed_after = set(), set()
    for rule in config["sheets"]:
        if not rule.get("enabled", True) or _skip_sheet(rule["name"], config):
            continue
        before, before_matches = _resolve_sheet(before_book, rule)
        after, after_matches = _resolve_sheet(after_book, rule)
        if any(_skip_sheet(sheet.name, config) for sheet in before_matches + after_matches):
            continue
        if len(before_matches) > 1 or len(after_matches) > 1:
            findings.add("LAYOUT_AMBIGUOUS", "More than one sheet matches the configured name/aliases. This sheet was not compared.",
                         sheet=rule["name"], baseline=len(before_matches) > 1)
            prepared.append((rule, None, None, None, None))
            continue
        if before:
            claimed_before.add(before.name)
        if after:
            claimed_after.add(after.name)
        if not before and rule.get("required", True):
            findings.add("TEMPLATE_ISSUE", "This required sheet is absent from the sent template. Check the configuration or the template before attributing omissions to the country.",
                         sheet=rule["name"], baseline=True)
        if before and after and before.name != after.name:
            findings.add("SHEET_RENAMED", f"Sheet {before.name!r} is named {after.name!r} in the received file. A configured alias allows comparison.", sheet=rule["name"])
        sent_layout = locate(before, rule) if before else None
        received_layout = locate(after, rule) if after else None
        prepared.append((rule, before, after, sent_layout, received_layout))
    return prepared, claimed_before, claimed_after


def _reference_map(prepared):
    result = {}
    for rule, sent, received, before, after in prepared:
        if not before or not after:
            continue
        old_groups, new_groups = _groups(before, rule["mode"]), _groups(after, rule["mode"])
        after_rows = {_row_key(row): row for row in after.rows}
        for row in before.rows:
            key = _name_group(row, rule["mode"])
            if len(old_groups[key]) != len(new_groups.get(key, [])) or row.ambiguous:
                continue
            target = after_rows.get(_row_key(row))
            if target is None or target.ambiguous or rule["mode"] == "identifiers" and len(old_groups[key]) != 1:
                continue
            for column, physical in row.columns.items():
                if column in target.columns:
                    result[(sent.name.casefold(), f"{get_column_letter(physical)}{row.row}")] = (
                        received.name, f"{get_column_letter(target.columns[column])}{target.row}")
    return result


def _unresolved_layout(layout):
    return layout is not None and any(problem["code"] in {"LAYOUT_UNRESOLVED", "AMBIGUOUS_LAYOUT"}
                                     for problem in layout.problems)


def _ambiguous_column(layout, row, period):
    return layout is not None and any(
        problem["code"] == "AMBIGUOUS_COLUMN" and normalize(problem.get("column")) == normalize(period)
        and (not problem.get("identifier") or normalize(problem["identifier"]) == normalize(row.identifier))
        for problem in layout.problems)


def _identifier_cells_equal(before, after):
    if numeric(before.value) and numeric(after.value):
        return before.value == after.value
    if isinstance(before.value, str) and isinstance(after.value, str):
        # Excel's usual exact lookups ignore case, but not extra whitespace.
        return before.value.casefold() == after.value.casefold()
    return type(before.value) is type(after.value) and before.value == after.value


def _table_inventory(sheet, layout, rule, findings, *, baseline):
    """Read actual headers of the one table proven by identifier layout discovery."""
    if sheet is None or layout is None or _unresolved_layout(layout):
        return [], None
    header_points = {range_boundaries(address)[:2] for address in layout.header_coordinates}
    identifier_columns = {row.label_column for row in layout.rows}
    if not identifier_columns:
        identifier_columns = {column for column, row in header_points
                              if normalize(sheet.cell(row, column).value) == normalize(rule.get("identifier_header", "KPI"))}
    if len(identifier_columns) != 1:
        return [], None
    identifier_column = next(iter(identifier_columns))
    candidates = []
    for table in sheet.tables:
        if rule.get("table_name") and normalize(rule["table_name"]) not in {normalize(table.get("name")), normalize(table.get("display_name"))}:
            continue
        left, top, right, bottom = range_boundaries(table["ref"])
        if (table.get("header_row_count", 1) != 1 or (identifier_column, top) not in header_points
                or not left <= identifier_column <= right):
            continue
        if all(left <= row.label_column <= right and top < row.row <= bottom for row in layout.rows):
            candidates.append((left, top, right))
    if len(candidates) != 1:
        return [], None
    left, top, right = candidates[0]
    entries, unavailable = [], []
    for column in range(left, right + 1):
        if _ignored(rule, top, column):
            continue
        value = sheet.cell(top, column)
        address = f"{get_column_letter(column)}{top}"
        if value.formula and not value.cache_available or value.kind in {"e", "b"}:
            unavailable.append(address)
        elif value.value is not None and str(value.value).strip():
            entries.append((str(value.value), address, column == identifier_column))
    names = [name for name, _, _ in entries]
    if unavailable:
        already_reported = {address for problem in layout.problems for address in problem.get("coordinates", [])}
        pending = [address for address in unavailable if address not in already_reported]
        if pending:
            findings.add("TEMPLATE_ISSUE" if baseline else "LAYOUT_AMBIGUOUS",
                         "Some structured-table column headers have no readable saved value; their presence cannot be compared reliably.",
                         sheet=rule["name"], baseline=baseline,
                         sent_cell=", ".join(pending) if baseline else "", received_cell="" if baseline else ", ".join(pending))
        return names, None
    return names, entries


def _table_column_checks(sent, received, before, after, rule, findings):
    old_names, old_entries = _table_inventory(sent, before, rule, findings, baseline=True)
    new_names, new_entries = _table_inventory(received, after, rule, findings, baseline=False)
    if old_entries is None or new_entries is None:
        return old_names, new_names
    configured = {normalize(name) for column in rule["columns"]
                  for name in [column["name"], *column.get("aliases", [])]}
    first, second = defaultdict(list), defaultdict(list)
    for entries, groups in ((old_entries, first), (new_entries, second)):
        for name, address, identifier in entries:
            if not identifier and normalize(name) not in configured:
                groups[normalize(name)].append((name, address))
    for key in sorted(first.keys() | second.keys()):
        old, new = first.get(key, []), second.get(key, [])
        if len(old) == len(new):
            continue
        added = len(new) > len(old)
        label = (new or old)[0][0]
        findings.add("COLUMN_ADDED" if added else "COLUMN_MISSING",
                     f"Structured-table header {label!r}: {len(old)} occurrence(s) in the sent table, {len(new)} in the received table. This column is outside the configured period rules.",
                     sheet=rule["name"], column=label,
                     sent_cell=", ".join(address for _, address in old), received_cell=", ".join(address for _, address in new),
                     before={"value": old_names, "formula": None}, after={"value": new_names, "formula": None})
    return old_names, new_names


def _sheet_result(rule, sent, received, before, after, findings, config, references):
    name, mode = rule["name"], rule["mode"]
    start = len(findings.items)
    _problems(before, findings, name, baseline=True, mode=mode)
    _problems(after, findings, name, sent_layout=before, mode=mode)
    table_before, table_after = _table_column_checks(sent, received, before, after, rule, findings) if mode == "identifiers" else ([], [])
    old_groups, new_groups, ambiguous = _integrity(before, after, rule, findings, name) if before and after else (_groups(before, mode), _groups(after, mode), set())
    if mode == "identifiers" and before and after:
        for key, old_rows in old_groups.items():
            new_rows = new_groups.get(key, [])
            if key in ambiguous or len(old_rows) != 1 or len(new_rows) != 1:
                continue
            old_row, new_row = old_rows[0], new_rows[0]
            if (_ignored(rule, old_row.row, old_row.label_column)
                    or _ignored(rule, new_row.row, new_row.label_column)):
                continue
            old_cell = sent.cell(old_row.row, old_row.label_column)
            new_cell = received.cell(new_row.row, new_row.label_column)
            if not _identifier_cells_equal(old_cell, new_cell):
                findings.add("KPI_INTEGRITY", "This identifier aligns after normalising its label, but its exact text or value type changed. Extra whitespace or numeric/text conversion can break exact lookups.",
                             sheet=name, kpi=old_row.identifier, column=rule.get("identifier_header", "KPI"),
                             sent_cell=f"{get_column_letter(old_row.label_column)}{old_row.row}",
                             received_cell=f"{get_column_letter(new_row.label_column)}{new_row.row}",
                             before=snapshot(old_cell), after=snapshot(new_cell))
    before_count, after_count = Counter(), Counter()
    after_rows = {_row_key(row): row for row in after.rows} if after else {}
    moves, seen = [], set()
    required_columns = {column["name"] for column in rule["columns"] if column.get("required", True)}
    configured_columns = {column["name"]: column for column in rule["columns"]}
    uncertain_layout = _unresolved_layout(after)
    for row in before.rows if before else []:
        group = _name_group(row, mode)
        target = after_rows.get(_row_key(row))
        for period, physical in row.columns.items():
            if period not in configured_columns or _ignored(rule, row.row, physical):
                continue
            slot = row.row, physical
            if slot in seen:
                continue
            seen.add(slot)
            old_cell = sent.cell(*slot)
            included = configured_columns[period].get("include_in_filling", True)
            new_physical = target.columns.get(period) if target else None
            if new_physical is not None and _ignored(rule, target.row, new_physical):
                continue
            if included:
                before_count["unknown" if row.ambiguous else state(old_cell)] += 1
            if (group in ambiguous or row.ambiguous or target is not None and target.ambiguous
                    or target is None and uncertain_layout or _ambiguous_column(after, target or row, period)):
                if included:
                    after_count["unknown"] += 1
                continue
            if new_physical is None:
                if included:
                    after_count["invalid"] += 1
                continue
            new_cell = received.cell(target.row, new_physical)
            if included:
                after_count[state(new_cell)] += 1
            old_address = f"{get_column_letter(physical)}{row.row}"
            new_address = f"{get_column_letter(new_physical)}{target.row}"
            if old_address != new_address:
                moves.append(f"{row.identifier} / {period}: {old_address} → {new_address}")
            context = dict(sheet=name, kpi=row.identifier, column=period, sent_cell=old_address, received_cell=new_address)
            _cell_checks(old_cell, new_cell, context, findings, config, rule, references, sent.name, received.name)
    if moves:
        findings.add("LAYOUT_CHANGED", f"{len(moves)} compared cell position(s) changed. Values were compared using their KPI identifiers and period labels. " + "\n".join(moves), sheet=name)
    missing = [rows[0].identifier for key, rows in old_groups.items()
               if len(rows) > len(new_groups.get(key, [])) and not (mode == "identifiers" and key in ambiguous and uncertain_layout)]
    unexpected = [rows[0].identifier for key, rows in new_groups.items()
                  if len(rows) > len(old_groups.get(key, [])) and not (mode == "identifiers" and key in ambiguous and _unresolved_layout(before))]
    notes = []
    if not before:
        notes.append("No confirmed sent-template cells are available for a filling comparison.")
    if _unresolved_layout(before) or before is not None and any(problem["code"] == "AMBIGUOUS_COLUMN" for problem in before.problems):
        notes.append("Some sent-template identifiers or headers could not be located reliably. Unlocated cells are excluded from the denominator; displayed coverage describes confirmed sent-template cells only.")
    if ambiguous or uncertain_layout:
        notes.append("Ambiguous locations are unknown, not empty. Values are not paired by guesswork.")
    if before_count["unknown"] or after_count["unknown"]:
        notes.append("An exact filling rate is unavailable where formula results are not saved or alignment is uncertain.")
    return {"name": name, "sent_name": sent.name if sent else None, "received_name": received.name if received else None,
            "detailed_checks": True, "is_kpi_sheet": normalize(name) == "kpi",
            "status": "matched" if sent and received else "missing" if sent else "added" if received else "absent",
            "expected_kpis": len(before.rows) if mode == "identifiers" and before else len(rule.get("kpis", [])),
            "sent_kpis": len(before.rows) if before else 0, "received_kpis": len(after.rows) if after else 0,
            "missing_kpis": missing, "unexpected_kpis": unexpected,
            "expected_columns": [column["name"] for column in rule["columns"] if column.get("required", True)],
            "sent_columns": _columns(before), "received_columns": _columns(after),
            "sent_table_columns": table_before, "received_table_columns": table_after,
            "missing_columns": sorted(required_columns - set(_columns(after))),
            "filling_before": filling(before_count), "filling_after": filling(after_count),
            "anomaly_count": len(findings.items) - start, "notes": notes}


def _skip_sheet(name, config):
    names = {name}
    for rule in config.get("sheets", []):
        group = [rule["name"], *rule.get("aliases", [])]
        if normalize(name) in {normalize(item) for item in group}:
            names.update(group)
    return any(fnmatchcase(candidate.casefold(), pattern.casefold())
               for candidate in names for pattern in config.get("ignore_sheets", []))


def _paired_formula_cells(prepared):
    """Yield only formulas with a proven, nonignored business-cell pairing."""
    for rule, sent, received, before, after in prepared:
        if before is None or after is None:
            continue
        old_groups, new_groups = _groups(before, rule["mode"]), _groups(after, rule["mode"])
        new_rows = {_row_key(row): row for row in after.rows}
        for row in before.rows:
            key = _name_group(row, rule["mode"])
            target = new_rows.get(_row_key(row))
            if (row.ambiguous or target is None or target.ambiguous
                    or len(old_groups[key]) != len(new_groups.get(key, []))
                    or rule["mode"] == "identifiers" and len(old_groups[key]) != 1):
                continue
            for period, physical in row.columns.items():
                new_column = target.columns.get(period)
                if (new_column is None or _ignored(rule, row.row, physical)
                        or _ignored(rule, target.row, new_column)):
                    continue
                old_cell, new_cell = sent.cell(row.row, physical), received.cell(target.row, new_column)
                if old_cell.formula and new_cell.formula:
                    yield sent, received, old_cell, new_cell, dict(
                        sheet=rule["name"], kpi=row.identifier, column=period,
                        sent_cell=f"{get_column_letter(physical)}{row.row}",
                        received_cell=f"{get_column_letter(new_column)}{target.row}")


def _formula_ranges(formula):
    from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
    try:
        return [token.value for token in Tokenizer(formula).items if token.subtype == "RANGE"]
    except (TokenizerError, ValueError, IndexError, TypeError):
        return []


def _resolve_name(book, token, sheet):
    """Respect local names overriding workbook names and exact sheet spelling."""
    names = {key.casefold(): (key, value) for key, value in book.defined_names.items()}
    qualifier, separator, name = token.rpartition("!")
    if separator:
        if qualifier.startswith("'") and qualifier.endswith("'"):
            qualifier = qualifier[1:-1].replace("''", "'")
        keys = [f"{qualifier}!{name}".casefold()]
    else:
        keys = [f"{sheet}!{token}".casefold(), token.casefold()]
    return next((names[key] for key in keys if key in names), None)


def _dependency_checks(prepared, before_book, after_book, findings, config, references):
    # Workbook metadata is examined only when a paired selected business cell
    # actually uses it. Helper-sheet links and unused names cannot create alerts.
    from .formulas import external_link_indices
    for sent, received, old_cell, new_cell, context in _paired_formula_cells(prepared):
        new_tokens = {token.casefold(): token for token in _formula_ranges(new_cell.formula)}
        for token in _formula_ranges(old_cell.formula):
            if token.casefold() not in new_tokens:
                continue  # The changed cell formula already supplies the evidence.
            old_name = _resolve_name(before_book, token, sent.name)
            new_name = _resolve_name(after_book, new_tokens[token.casefold()], received.name)
            if old_name is None:
                continue
            name, old_reference = old_name
            new_reference = new_name[1] if new_name else None
            old_formula = old_reference if old_reference.startswith("=") else "=" + old_reference
            new_formula = (new_reference if new_reference.startswith("=") else "=" + new_reference) if new_reference is not None else None
            if new_formula is None or not same_formula(old_formula, new_formula, sent.name, received.name, references,
                                                       ignore_case=config.get("ignore_formula_case", True)):
                findings.add("DEFINED_NAME_CHANGED", f"Name {name!r}, used by this formula, was removed or points to a different location.",
                             **context, before={**snapshot(old_cell), "value": old_reference},
                             after={**snapshot(new_cell), "value": new_reference})
        old_indices = external_link_indices(old_cell.formula)
        new_indices = external_link_indices(new_cell.formula)
        for index in sorted(old_indices & new_indices):
            old_target = getattr(before_book, "external_link_targets", {}).get(index)
            new_target = getattr(after_book, "external_link_targets", {}).get(index)
            if old_target is not None and new_target != old_target:
                findings.add("EXTERNAL_LINK_CHANGED", "The external workbook target used by this formula changed or was removed. The target file's availability is not tested.",
                             **context, before={**snapshot(old_cell), "value": old_target},
                             after={**snapshot(new_cell), "value": new_target})


def compare_pair(before_book, after_book, config, country):
    findings = Findings(config, country)
    prepared, claimed_before, claimed_after = _prepare_sheets(before_book, after_book, config, findings)
    matched_raw = {name.casefold() for name in before_book.sheetnames} & {name.casefold() for name in after_book.sheetnames}
    for name in before_book.sheetnames:
        if name.casefold() not in matched_raw and name not in claimed_before.intersection(claimed_after):
            if not any(sent and received and sent.name == name for _, sent, received, _, _ in prepared):
                important = any(sent and sent.name == name and rule.get("required", True)
                                for rule, sent, _, _, _ in prepared)
                findings.add("SHEET_MISSING", "A sheet present in the sent workbook is absent from the received workbook.",
                             sheet=name, severity="high" if important else "medium")
    for name in after_book.sheetnames:
        if name.casefold() not in matched_raw:
            if not any(sent and received and received.name == name for _, sent, received, _, _ in prepared):
                findings.add("SHEET_ADDED", "This sheet was added in the received workbook.", sheet=name)
    references = _reference_map(prepared)
    sheets = [_sheet_result(*item, findings, config, references) for item in prepared]
    known_names = {name for sheet in sheets for name in (sheet["sent_name"], sheet["received_name"]) if name}
    for name in dict.fromkeys(before_book.sheetnames + after_book.sheetnames):
        if name in known_names:
            continue
        sheets.append({"name": name, "sent_name": name if name in before_book.sheets else None,
                       "received_name": name if name in after_book.sheets else None,
                       "status": "inventory_only", "detailed_checks": False, "is_kpi_sheet": normalize(name) == "kpi",
                       "expected_kpis": None, "sent_kpis": None, "received_kpis": None,
                       "missing_kpis": [], "unexpected_kpis": [], "expected_columns": [], "sent_columns": [],
                       "received_columns": [], "missing_columns": [], "filling_before": filling(), "filling_after": filling(),
                       "anomaly_count": 0, "notes": ["Sheet presence only; its contents are outside the selected checks."]})
    _dependency_checks(prepared, before_book, after_book, findings, config, references)
    selected = [sheet for sheet in sheets if sheet["detailed_checks"] and not sheet["is_kpi_sheet"]]
    kpi_sheets = [sheet for sheet in sheets if sheet["detailed_checks"] and sheet["is_kpi_sheet"]]
    for sheet in sheets:
        sheet["anomaly_count"] = sum(finding["sheet"] in {sheet["name"], sheet.get("sent_name"), sheet.get("received_name")} for finding in findings.items)
    country.update(sheets=sheets, filling_before=sum_filling([item["filling_before"] for item in selected]),
                   filling_after=sum_filling([item["filling_after"] for item in selected]),
                   kpi_filling_before=sum_filling([item["filling_before"] for item in kpi_sheets]),
                   kpi_filling_after=sum_filling([item["filling_after"] for item in kpi_sheets]),
                   anomaly_count=len(findings.items), by_severity=dict(Counter(item["severity"] for item in findings.items)),
                   sent_sheet_count=len(before_book.sheetnames), received_sheet_count=len(after_book.sheetnames),
                   notes=["Overall filling covers selected business sheets and excludes KPI. Missing received locations remain in the denominator; additional received rows cannot inflate the rate."])
    if not selected or not country["filling_before"]["expected"]:
        country["notes"].append("No cells could be measured for overall filling. KPI is measured separately.")
    return country, findings.items


def _file_key(path):
    return unicodedata.normalize("NFC", path.name).casefold()


def discover(directory):
    result = defaultdict(list)
    ignored = []
    for path in sorted(Path(directory).iterdir(), key=lambda value: value.name.casefold()):
        if path.is_file() and not path.name.startswith("~$") and path.suffix.casefold() in {".xlsx", ".xlsm", ".xls", ".xlsb"}:
            result[_file_key(path)].append(path)
        elif path.is_file():
            ignored.append(path.name)
    return result, ignored


def analyse(config, *, progress=None):
    report = progress or (lambda message: None)
    for name in ("sent_dir", "received_dir"):
        if not Path(config[name]).is_dir():
            raise ValueError(f"Folder not found: {config[name]}. Create it and add the country workbooks.")
    sent, ignored_sent = discover(config["sent_dir"])
    received, ignored_received = discover(config["received_dir"])
    countries, anomalies, notices = [], [], []
    stats = Counter(sent_files=sum(map(len, sent.values())), received_files=sum(map(len, received.values())))
    for key in sorted(sent.keys() | received.keys()):
        old, new = sent.get(key, []), received.get(key, [])
        representative = (old or new)[0]
        country = {"id": fingerprint(["country", key]), "name": config.get("country_names", {}).get(key, representative.stem),
                   "file_key": key, "sent_file": old[0].name if len(old) == 1 else None,
                   "received_file": new[0].name if len(new) == 1 else None,
                   "status": "matched", "sheets": [], "notes": [], "anomaly_count": 0, "by_severity": {},
                   "filling_before": filling(), "filling_after": filling(),
                   "kpi_filling_before": filling(), "kpi_filling_after": filling()}
        findings = Findings(config, country)
        if len(old) > 1 or len(new) > 1:
            country["status"] = "ambiguous"
            stats["ambiguous_pairs"] += 1
            findings.add("FILE_AMBIGUOUS", "Multiple files have the same case-insensitive filename. No pair was guessed. " + ", ".join(path.name for path in old + new))
        elif not new:
            country["status"] = "not_received"
            stats["missing_received"] += 1
            findings.add("FILE_MISSING", "No received workbook has this exact filename. Rename files only when you have confirmed the country pairing.")
        elif not old:
            country["status"] = "unexpected"
            stats["unexpected_received"] += 1
            findings.add("FILE_UNEXPECTED", "This received filename has no sent counterpart. It is excluded from filling-rate comparisons.")
        else:
            stats["matched_pairs"] += 1
            report(f"Comparing {representative.name}...")
            try:
                selected_names = [name for rule in config["sheets"] if rule.get("enabled", True)
                                  and not _skip_sheet(rule["name"], config)
                                  for name in [rule["name"], *rule.get("aliases", [])]]
                before_book = read_workbook(old[0], selected_sheets=selected_names)
                after_book = read_workbook(new[0], selected_sheets=selected_names)
            except Exception as exc:
                country["status"] = "read_error"
                stats["read_errors"] += 1
                findings.add("FILE_READ_ERROR", f"This pair could not be analysed: {type(exc).__name__}: {exc}. Use readable, unencrypted .xlsx or .xlsm files. Other countries were still processed.")
            else:
                # Comparison failures are program/configuration failures, not
                # evidence that a country's workbook is corrupt.
                country, items = compare_pair(before_book, after_book, config, country)
                findings.items.extend(items)
                stats["analysed_pairs"] += 1
        visible = [item for item in findings.items if not item["baseline"]]
        country["anomaly_count"] = len(visible)
        country["by_severity"] = dict(Counter(item["severity"] for item in visible))
        countries.append(country)
        anomalies.extend(findings.items)
    for key in ("sent_files", "received_files", "matched_pairs", "analysed_pairs", "missing_received", "unexpected_received", "ambiguous_pairs", "read_errors"):
        stats.setdefault(key, 0)
    measured = [country for country in countries if country["status"] == "matched"]
    summary = dict(stats)
    visible = [item for item in anomalies if not item["baseline"]]
    summary.update(countries_with_issues=sum(any(item["country_id"] == country["id"] for item in visible) for country in countries),
                   anomaly_count=len(visible),
                   by_severity=dict(Counter(item["severity"] for item in visible)),
                   filling_before=sum_filling([item["filling_before"] for item in measured]),
                   filling_after=sum_filling([item["filling_after"] for item in measured]),
                   kpi_filling_before=sum_filling([item["kpi_filling_before"] for item in measured]),
                   kpi_filling_after=sum_filling([item["kpi_filling_after"] for item in measured]))
    notices.append("Workbooks are read without being changed or recalculated. Formula filling uses saved Excel results; unavailable results are unknown. Save recalculated workbooks in Excel before checking if those results are needed.")
    notices.append("Overall filling covers only the selected business sheets, excluding KPI. The KPI sheet has its own separate filling rate. Other sheets are checked only for addition or removal.")
    notices.append("Filling = a numeric, nonzero value. Negative numbers and numeric percentages count; blanks, zeros, text and booleans do not. Percentages are weighted by expected cell counts.")
    if ignored_sent or ignored_received:
        notices.append(f"Ignored temporary/non-workbook files: {len(ignored_sent)} in sent; {len(ignored_received)} in received.")
    if not countries:
        notices.append("No country Excel files were found. Add the sent templates and received returns to their configured folders, then run the report again.")
    return {"schema_version": 1, "report_id": fingerprint([str(config["sent_dir"]), str(config["received_dir"])]),
            "generated_at": datetime.now(timezone.utc).isoformat(), "title": config.get("report_title", "POPS return review"),
            "demo": False, "settings": {key: config.get(key) for key in ("filling_scope", "historical_columns", "prefilled_changes", "checks", "severity_overrides")},
            "summary": summary, "countries": countries, "anomalies": anomalies, "notices": notices}
