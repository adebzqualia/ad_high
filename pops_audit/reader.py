"""Read sparse workbook values and formula caches without loading Excel styles.

This module never writes a workbook or evaluates a formula. A missing cached
formula result remains unknown, distinct from a genuinely empty input cell.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import posixpath
import re
import unicodedata
from urllib.parse import unquote
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.formula.tokenizer import TokenizerError


@dataclass(frozen=True)
class Cell:
    value: object = None
    formula: str | None = None
    kind: str = "n"
    cache_available: bool = True


@dataclass
class Sheet:
    name: str
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)
    merges: list[str] = field(default_factory=list)
    tables: list[dict] = field(default_factory=list)
    sheet_type: str = "worksheet"

    def cell(self, row: int, column: int) -> Cell:
        if not 1 <= row <= 1048576 or not 1 <= column <= 16384:
            raise ValueError("Cell coordinate exceeds Excel worksheet bounds")
        return self.cells.get((row, column), Cell())


@dataclass
class Book:
    sheets: dict[str, Sheet] = field(default_factory=dict)
    sheetnames: list[str] = field(default_factory=list)
    defined_names: dict[str, str] = field(default_factory=dict)
    external_links: list[str] = field(default_factory=list)
    external_link_targets: dict[int, str] = field(default_factory=dict)


def coordinate(row: int, column: int) -> str:
    if not 1 <= row <= 1048576 or not 1 <= column <= 16384:
        raise ValueError("Cell coordinate exceeds Excel worksheet bounds")
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def position(reference: str) -> tuple[int, int]:
    match = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]{0,6})", str(reference))
    if not match:
        raise ValueError(f"Invalid Excel cell reference: {reference!r}")
    column = 0
    for letter in match.group(1).upper():
        column = column * 26 + ord(letter) - 64
    row = int(match.group(2))
    coordinate(row, column)
    return row, column


def bounds(reference: str) -> tuple[int, int, int, int]:
    parts = str(reference).split(":")
    if len(parts) not in (1, 2):
        raise ValueError(f"Invalid Excel range: {reference!r}")
    top, left = position(parts[0])
    bottom, right = position(parts[-1])
    if bottom < top or right < left:
        raise ValueError(f"Invalid Excel range: {reference!r}")
    return top, left, bottom, right


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _child(element, name):
    return next((child for child in element if _local(child.tag) == name), None)


def _relationship_id(element):
    return next((value for key, value in element.attrib.items()
                 if key.startswith("{") and _local(key) == "id"), None)


def _resolve(part, target):
    target = unquote(target).replace("\\", "/")
    resolved = posixpath.normpath(target.lstrip("/") if target.startswith("/")
                                 else posixpath.join(posixpath.dirname(part), target))
    if resolved in {".", ".."} or resolved.startswith("../"):
        raise ValueError(f"Workbook relationship escapes package: {target!r}")
    return resolved


def _relationships(archive, part):
    folder, name = posixpath.split(part)
    path = posixpath.join(folder, "_rels", name + ".rels") if part else "_rels/.rels"
    if path not in archive.namelist():
        return {}
    return {item.get("Id"): item.attrib for item in ET.fromstring(archive.read(path))
            if _local(item.tag) == "Relationship"}


def _text(container):
    if container is None:
        return ""
    result = []
    for item in container:
        if _local(item.tag) == "t":
            result.append(item.text or "")
        elif _local(item.tag) == "r":
            result.extend(child.text or "" for child in item if _local(child.tag) == "t")
    return "".join(result)


def _value(node, strings, has_formula):
    kind = node.get("t", "n")
    cached = _child(node, "v")
    raw = cached.text if cached is not None else None
    if kind == "inlineStr":
        inline = _child(node, "is")
        return _text(inline), "s", inline is not None or not has_formula
    if raw is None:
        if kind == "str" and cached is not None:
            return "", "s", True
        return None, "s" if kind in {"s", "str"} else kind, not has_formula
    if kind == "s":
        index = int(raw)
        if not 0 <= index < len(strings):
            raise ValueError(f"Invalid shared string index: {index}")
        return strings[index], "s", True
    if kind == "b":
        if raw.casefold() not in {"1", "0", "true", "false"}:
            raise ValueError(f"Invalid Boolean cell value: {raw!r}")
        return raw.casefold() in {"1", "true"}, "b", True
    if kind in {"e", "str", "d"}:
        return raw, "s" if kind == "str" else kind, True
    if kind not in {"n", ""}:
        raise ValueError(f"Unsupported cell type: {kind!r}")
    return int(raw) if re.fullmatch(r"[+-]?[0-9]+", raw) else float(raw), "n", True


def _read_sheet(archive, part, sheet, strings):
    tree = ET.fromstring(archive.read(part))
    merged = _child(tree, "mergeCells")
    if merged is not None:
        sheet.merges = [item.get("ref") for item in merged if item.get("ref")]
        for reference in sheet.merges:
            bounds(reference)
    records, masters = [], {}
    data = _child(tree, "sheetData")
    previous_row = 0
    for row_node in data if data is not None else ():
        row = int(row_node.get("r", previous_row + 1))
        previous_row, previous_column = row, 0
        for node in row_node:
            if _local(node.tag) != "c":
                continue
            cell_row, column = position(node.get("r")) if node.get("r") else (row, previous_column + 1)
            reference = coordinate(cell_row, column)
            previous_column = column
            formula = _child(node, "f")
            records.append((cell_row, column, reference, node, formula))
            if formula is not None and formula.get("t") == "shared" and formula.text:
                masters[formula.get("si")] = reference, formula.text
    for row, column, reference, node, formula in records:
        formula_text = None
        if formula is not None:
            formula_text = formula.text
            if formula.get("t") == "shared" and not formula_text:
                master = masters.get(formula.get("si"))
                if master is None:
                    raise ValueError(f"Missing shared-formula master at {sheet.name}!{reference}")
                try:
                    formula_text = Translator("=" + master[1].lstrip("="), origin=master[0]).translate_formula(reference)
                except (ValueError, TypeError, IndexError, TranslatorError, TokenizerError) as exc:
                    raise ValueError(f"Cannot expand shared formula at {sheet.name}!{reference}: {exc}") from exc
            formula_text = "=" + (formula_text or "").lstrip("=")
        try:
            value, kind, available = _value(node, strings, formula is not None)
        except (ValueError, TypeError, IndexError) as exc:
            raise ValueError(f"Cannot read {sheet.name}!{reference}: {exc}") from exc
        # Ignore style-only cells and inflated worksheet dimensions entirely.
        if value is None and formula_text is None:
            continue
        sheet.cells[row, column] = Cell(value, formula_text, kind, available)
    relationships = _relationships(archive, part)
    table_parts = _child(tree, "tableParts")
    for item in table_parts if table_parts is not None else ():
        relation = relationships.get(_relationship_id(item))
        if relation is None or relation.get("TargetMode") == "External":
            raise ValueError(f"Missing or external table relationship in {sheet.name}")
        table = ET.fromstring(archive.read(_resolve(part, relation.get("Target", ""))))
        reference = table.get("ref", "")
        bounds(reference)
        columns = _child(table, "tableColumns")
        sheet.tables.append({
            "name": table.get("name", table.get("displayName", "")),
            "display_name": table.get("displayName", table.get("name", "")),
            "ref": reference, "header_row_count": int(table.get("headerRowCount", "1")),
            "totals_row_count": int(table.get("totalsRowCount", "0")),
            "columns": [{"id": column.get("id"), "name": column.get("name", "")}
                        for column in columns if _local(column.tag) == "tableColumn"] if columns is not None else [],
        })


def read_workbook(path: str | Path, selected_sheets=None) -> Book:
    """Read selected sheets' cells while retaining the complete sheet inventory.

    ``selected_sheets`` accepts configured names and aliases. Its controlled
    presentation normalization matches the sheet resolver; formula references
    independently preserve the exact identity of Excel sheet names.
    """
    result = Book()
    selected = None if selected_sheets is None else {
        " ".join(unicodedata.normalize("NFKC", str(name)).split()).casefold()
        for name in selected_sheets
    }
    with ZipFile(path) as archive:
        if len(archive.namelist()) != len(set(archive.namelist())):
            raise ValueError("Workbook has duplicate ZIP part names")
        roots = _relationships(archive, "")
        document = next((item for item in roots.values()
                         if item.get("Type", "").endswith("/officeDocument")
                         and item.get("TargetMode") != "External"), None)
        workbook_part = _resolve("", document.get("Target", "")) if document else "xl/workbook.xml"
        workbook = ET.fromstring(archive.read(workbook_part))
        relationships = _relationships(archive, workbook_part)
        string_relation = next((item for item in relationships.values()
                                if item.get("Type", "").endswith("/sharedStrings")
                                and item.get("TargetMode") != "External"), None)
        strings = []
        if string_relation:
            strings = [_text(item) for item in ET.fromstring(archive.read(_resolve(workbook_part, string_relation["Target"])))
                       if _local(item.tag) == "si"]
        sheets = _child(workbook, "sheets")
        for item in sheets if sheets is not None else ():
            name = item.get("name", "")
            if name.casefold() in {existing.casefold() for existing in result.sheetnames}:
                raise ValueError(f"Duplicate worksheet name: {name!r}")
            relation = relationships.get(_relationship_id(item))
            if relation is None or relation.get("TargetMode") == "External":
                raise ValueError(f"Missing worksheet relationship: {name!r}")
            kind = relation.get("Type", "").rsplit("/", 1)[-1]
            sheet = Sheet(name, sheet_type=kind)
            result.sheets[name] = sheet
            result.sheetnames.append(name)
            selection_key = " ".join(unicodedata.normalize("NFKC", name).split()).casefold()
            if kind == "worksheet" and (selected is None or selection_key in selected):
                _read_sheet(archive, _resolve(workbook_part, relation.get("Target", "")), sheet, strings)
        names = _child(workbook, "definedNames")
        for item in names if names is not None else ():
            name = item.get("name", "")
            scope = item.get("localSheetId")
            if scope is not None:
                index = int(scope)
                name = f"{result.sheetnames[index] if 0 <= index < len(result.sheetnames) else scope}!{name}"
            result.defined_names[name] = item.text or ""
        targets_by_relationship = {}
        for relationship_id, relation in relationships.items():
            if not relation.get("Type", "").endswith("/externalLink"):
                continue
            part = _resolve(workbook_part, relation.get("Target", ""))
            targets = []
            for link in _relationships(archive, part).values():
                if (link.get("TargetMode") == "External" and link.get("Target")
                        and link.get("Type", "").endswith("/externalLinkPath")):
                    target = unquote(link["Target"])
                    result.external_links.append(target)
                    targets.append(target)
            if len(set(targets)) == 1:
                targets_by_relationship[relationship_id] = targets[0]
        external_references = _child(workbook, "externalReferences")
        for index, item in enumerate(external_references if external_references is not None else (), 1):
            target = targets_by_relationship.get(_relationship_id(item))
            if target is not None:
                result.external_link_targets[index] = target
    result.external_links = list(dict.fromkeys(result.external_links))
    return result
